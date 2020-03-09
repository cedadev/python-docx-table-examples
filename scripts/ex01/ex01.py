from docx import Document

doc = Document('Doc1.docx')
table = doc.tables[0]

table.cell(1, 0).text = str('blahblah')
doc.save('Doc1.docx')
